# core/datos.py
"""
DatosJuego: motor que interpreta Black_Seas.xlsx como reglamento digital.
Provee API para construir Barco base (BARCOS + ARMAS + NORMAS) y para
recuperar MEJORAS, OFICIALES, DAÑOS por identificador/nombre.
"""

from pathlib import Path
import logging
from typing import Optional, Dict, Any, List, Iterable

logger = logging.getLogger(__name__)
logging.basicConfig(level=logging.INFO)

PROJECT_ROOT = Path(__file__).resolve().parent.parent
XLSX_DEFAULT = PROJECT_ROOT / "data" / "Black_Seas.xlsx"

# Prefer pandas if está disponible; fallback a openpyxl
try:
    import pandas as pd  # type: ignore
    PANDAS = True
except Exception:
    pd = None  # type: ignore
    PANDAS = False

try:
    from openpyxl import load_workbook  # type: ignore
    OPENPYXL = True
except Exception:
    load_workbook = None  # type: ignore
    OPENPYXL = False

def _read_sheet_pandas(path: Path, sheet: str):
    try:
        df = pd.read_excel(path, sheet_name=sheet, engine="openpyxl")
        df.columns = [str(c).strip() for c in df.columns]
        return df
    except Exception:
        return None

def _read_sheet_openpyxl(path: Path, sheet: str):
    try:
        wb = load_workbook(path, data_only=True)
        if sheet not in wb.sheetnames:
            return None
        ws = wb[sheet]
        rows = list(ws.values)
        if not rows:
            return None
        headers = [str(h).strip() if h is not None else "" for h in rows[0]]
        records = []
        for r in rows[1:]:
            rec = {}
            for k, v in zip(headers, r):
                rec[k] = v
            records.append(rec)
        return records
    except Exception:
        return None

class DatosJuego:
    def __init__(self, xlsx_path: Optional[Path] = None):
        self.xlsx_path = Path(xlsx_path) if xlsx_path else XLSX_DEFAULT
        self._cache: Dict[str, Any] = {}
        self._load_all()

    def _load_sheet(self, name: str):
        if not self.xlsx_path.exists():
            logger.error("Excel no encontrado: %s", self.xlsx_path)
            return None
        if PANDAS:
            df = _read_sheet_pandas(self.xlsx_path, name)
            if df is not None:
                return df
        if OPENPYXL:
            recs = _read_sheet_openpyxl(self.xlsx_path, name)
            return recs
        return None

    def _load_all(self):
        # Cargar hojas principales en cache
        for sheet in ("BARCOS", "ARMAS", "OFICIALES", "DAÑOS", "NORMAS", "MEJORAS"):
            try:
                self._cache[sheet] = self._load_sheet(sheet)
                logger.info("Hoja cargada: %s -> %s", sheet, type(self._cache[sheet]).__name__)
            except Exception:
                logger.exception("Error cargando hoja %s", sheet)
                self._cache[sheet] = None

    # ---------- utilidades ----------
    def _df_to_records(self, df):
        if df is None:
            return []
        if PANDAS and hasattr(df, "to_dict"):
            return df.fillna("").to_dict(orient="records")
        return list(df)  # openpyxl fallback already list of dicts

    # ---------- API pública ----------
    def list_barcos(self) -> List[Dict[str, Any]]:
        return self._df_to_records(self._cache.get("BARCOS"))

    def get_barco_base_by_id(self, identifier: Any) -> Optional[Dict[str, Any]]:
        """
        identifier puede ser ID interno (si existe columna ID) o Nombre.
        Devuelve el registro base de BARCOS (sin armas ni normas aplicadas).
        """
        df = self._cache.get("BARCOS")
        if df is None:
            return None
        # si pandas DataFrame
        try:
            if PANDAS and hasattr(df, "iterrows"):
                cols = df.columns
                id_col = next((c for c in cols if str(c).strip().lower() in ("id", "id_inter", "id_interno")), None)
                name_col = next((c for c in cols if "nom" in str(c).lower()), None)
                if id_col and identifier is not None and str(identifier).strip().isdigit():
                    res = df[df[id_col].astype(str).str.strip() == str(identifier).strip()]
                    if not res.empty:
                        return res.iloc[0].to_dict()
                if name_col:
                    res = df[df[name_col].astype(str).str.strip() == str(identifier).strip()]
                    if not res.empty:
                        return res.iloc[0].to_dict()
                # fallback: search any column for match
                for _, r in df.iterrows():
                    if any(str(v).strip() == str(identifier).strip() for v in r.values):
                        return r.to_dict()
                return None
            else:
                # openpyxl list of dicts
                for r in df:
                    for v in r.values():
                        if v is None:
                            continue
                        if str(v).strip() == str(identifier).strip():
                            return r
                return None
        except Exception:
            logger.exception("Error buscando barco por id/nombre")
            return None

    def get_armas_for_barco(self, barco_record: Dict[str, Any]) -> List[Dict[str, Any]]:
        """
        Devuelve lista de registros de ARMAS asociados al barco (por Nombre, Tipo, Facción).
        ARMAS debe contener columnas identificadoras (Nombre, Tipo, Facción).
        """
        df = self._cache.get("ARMAS")
        if df is None:
            return []
        try:
            nombre = (barco_record.get("NOMBRE") or barco_record.get("Nombre") or "").strip()
            tipo = (barco_record.get("TIPO") or barco_record.get("Tipo") or "").strip()
            fac = (barco_record.get("Facción") or barco_record.get("FACCION") or barco_record.get("Faccion") or "").strip()
            if PANDAS and hasattr(df, "iterrows"):
                cols = df.columns
                col_n = next((c for c in cols if "nom" in str(c).lower()), None)
                col_t = next((c for c in cols if "tipo" in str(c).lower()), None)
                col_f = next((c for c in cols if "facc" in str(c).lower() or "facción" in str(c).lower()), None)
                mask = None
                if col_n and nombre:
                    mask = df[col_n].astype(str).str.strip() == nombre
                if col_t and tipo:
                    m2 = df[col_t].astype(str).str.strip() == tipo
                    mask = m2 if mask is None else (mask & m2)
                if col_f and fac:
                    m3 = df[col_f].astype(str).str.strip() == fac
                    mask = m3 if mask is None else (mask & m3)
                if mask is None:
                    return df.to_dict(orient="records")
                res = df[mask]
                return res.to_dict(orient="records")
            else:
                # openpyxl fallback: filter list of dicts
                recs = []
                for r in df:
                    match = True
                    if nombre:
                        # try match by any name-like field
                        if not any("nom" in str(k).lower() and str(r.get(k) or "").strip() == nombre for k in r.keys()):
                            match = False
                    if tipo and match:
                        if not any("tipo" in str(k).lower() and str(r.get(k) or "").strip() == tipo for k in r.keys()):
                            match = False
                    if fac and match:
                        if not any(("facc" in str(k).lower() or "facción" in str(k).lower()) and str(r.get(k) or "").strip() == fac for k in r.keys()):
                            match = False
                    if match:
                        recs.append(r)
                return recs
        except Exception:
            logger.exception("Error obteniendo armas para barco")
            return []

    def get_oficial_by_name_or_id(self, identifier: Any) -> Optional[Dict[str, Any]]:
        df = self._cache.get("OFICIALES")
        if df is None:
            return None
        try:
            if PANDAS and hasattr(df, "iterrows"):
                cols = df.columns
                id_col = next((c for c in cols if "id" in str(c).lower()), None)
                name_col = next((c for c in cols if "nom" in str(c).lower()), None)
                if id_col and identifier is not None and str(identifier).strip().isdigit():
                    res = df[df[id_col].astype(str).str.strip() == str(identifier).strip()]
                    if not res.empty:
                        return res.iloc[0].to_dict()
                if name_col:
                    res = df[df[name_col].astype(str).str.strip() == str(identifier).strip()]
                    if not res.empty:
                        return res.iloc[0].to_dict()
                # fallback
                for _, r in df.iterrows():
                    if any(str(v).strip() == str(identifier).strip() for v in r.values):
                        return r.to_dict()
                return None
            else:
                for r in df:
                    if any(str(v).strip() == str(identifier).strip() for v in r.values):
                        return r
                return None
        except Exception:
            logger.exception("Error buscando oficial")
            return None

    def get_mejora_by_name(self, name: str) -> Optional[Dict[str, Any]]:
        df = self._cache.get("MEJORAS")
        if df is None:
            return None
        try:
            if PANDAS and hasattr(df, "iterrows"):
                cols = df.columns
                col_n = next((c for c in cols if "nom" in str(c).lower() or "name" in str(c).lower()), None)
                if col_n:
                    res = df[df[col_n].astype(str).str.strip() == str(name).strip()]
                    if not res.empty:
                        return res.iloc[0].to_dict()
                # fallback search
                for _, r in df.iterrows():
                    if any(str(v).strip() == str(name).strip() for v in r.values):
                        return r.to_dict()
                return None
            else:
                for r in df:
                    if any(str(v).strip() == str(name).strip() for v in r.values):
                        return r
                return None
        except Exception:
            logger.exception("Error buscando mejora")
            return None

    def get_danos(self) -> List[Dict[str, Any]]:
        return self._df_to_records(self._cache.get("DAÑOS"))

    def get_norma_por_faccion(self, faccion: str) -> Optional[Dict[str, Any]]:
        df = self._cache.get("NORMAS")
        if df is None:
            return None
        try:
            if PANDAS and hasattr(df, "iterrows"):
                cols = df.columns
                col_f = next((c for c in cols if "facc" in str(c).lower() or "facción" in str(c).lower()), None)
                if col_f:
                    res = df[df[col_f].astype(str).str.strip() == str(faccion).strip()]
                    if not res.empty:
                        return res.iloc[0].to_dict()
                # fallback: search any column
                for _, r in df.iterrows():
                    if any(str(v).strip() == str(faccion).strip() for v in r.values):
                        return r.to_dict()
                return None
            else:
                for r in df:
                    if any(str(v).strip() == str(faccion).strip() for v in r.values):
                        return r
                return None
        except Exception:
            logger.exception("Error buscando norma por facción")
            return None

    # Método utilitario para reconstruir "barco base" combinando BARCOS + ARMAS + NORMAS
    def build_barco_base(self, barco_identifier: Any) -> Optional[Dict[str, Any]]:
        """
        Devuelve un dict que representa el BARCO BASE:
        - datos de BARCOS (registro)
        - armas: lista de registros ARMAS
        - norma_nacional: registro de NORMAS para su facción (si existe)
        """
        base = self.get_barco_base_by_id(barco_identifier)
        if not base:
            return None
        armas = self.get_armas_for_barco(base)
        norma = self.get_norma_por_faccion(base.get("Facción") or base.get("FACCION") or "")
        # Normalizar claves: devolver copia para no mutar cache
        result = dict(base)
        result["_ARMAS"] = armas
        result["_NORMA"] = norma
        return result
